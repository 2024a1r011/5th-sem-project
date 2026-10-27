"""Generate sample PDF/DOCX resumes from data/source_text/*.txt.

Real resumes come in many layouts, so we create a few variants to test the
extractor: plain PDF, DOCX, and a DOCX whose content sits inside a table.

Run from the project root:  python scripts/generate_samples.py
"""
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "source_text"
OUT = ROOT / "data" / "sample_resumes"
HEADINGS = {"SUMMARY", "OBJECTIVE", "EDUCATION", "EXPERIENCE", "PROJECTS",
            "SKILLS", "CERTIFICATIONS", "ACHIEVEMENTS"}


def esc(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def make_pdf(lines, out_path):
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=10, leading=13)
    head = ParagraphStyle("head", parent=styles["Heading2"], fontSize=12,
                          spaceBefore=8, spaceAfter=2)
    name = ParagraphStyle("name", parent=styles["Title"], fontSize=18)
    bullet = ParagraphStyle("bullet", parent=body, leftIndent=14, bulletIndent=4)
    doc = SimpleDocTemplate(str(out_path), pagesize=A4,
                            leftMargin=40, rightMargin=40, topMargin=36, bottomMargin=36)
    story = []
    for i, line in enumerate(lines):
        if not line.strip():
            story.append(Spacer(1, 4))
        elif i == 0:
            story.append(Paragraph(esc(line), name))
        elif line.strip() in HEADINGS:
            story.append(Paragraph(esc(line.strip().title()), head))
        elif line.startswith("- "):
            story.append(Paragraph(esc(line[2:]), bullet, bulletText="\u2022"))
        else:
            story.append(Paragraph(esc(line), body))
    doc.build(story)


def make_docx(lines, out_path, in_table=False):
    doc = Document()
    target = doc
    if in_table:                      # content inside a 1x1 table
        table = doc.add_table(rows=1, cols=1)
        target = table.cell(0, 0)
        target.paragraphs[0].text = ""
    for i, line in enumerate(lines):
        if not line.strip():
            continue
        if i == 0:
            p = target.add_paragraph(); p.add_run(line).bold = True
        elif line.strip() in HEADINGS:
            p = target.add_paragraph(); p.add_run(line.strip().title()).bold = True
        elif line.startswith("- "):
            target.add_paragraph("\u2022 " + line[2:])
        else:
            target.add_paragraph(line)
    doc.save(str(out_path))


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for txt in sorted(SRC.glob("*.txt")):
        lines = txt.read_text(encoding="utf-8").splitlines()
        stem = txt.stem
        make_pdf(lines, OUT / f"{stem}.pdf")
        print("wrote", OUT / f"{stem}.pdf")
    # DOCX variants
    a = (SRC / "resume_ananya.txt").read_text(encoding="utf-8").splitlines()
    r = (SRC / "resume_rohan.txt").read_text(encoding="utf-8").splitlines()
    make_docx(a, OUT / "resume_ananya.docx")
    make_docx(r, OUT / "resume_rohan_table.docx", in_table=True)
    print("wrote DOCX variants")


if __name__ == "__main__":
    main()
