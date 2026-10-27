"""Step 1: Resume input -> clean text.

Supports PDF, DOCX and TXT. Returns an `ExtractionResult` with the cleaned
text plus warnings (e.g. "no text layer, probably a scanned PDF").

CLI:
    python -m resume_analyzer.extraction path/to/resume.pdf
"""
from __future__ import annotations

import re
import sys
import unicodedata
from pathlib import Path
from typing import List, Tuple

from .schema import ExtractionResult

SUPPORTED = {".pdf", ".docx", ".txt"}

# Characters that commonly act as bullets in resumes.
_BULLETS = "•▪●◦■□‣⁃∙·◆◇➢➤►▸–—*"
_BULLET_RE = re.compile(rf"^\s*[{re.escape(_BULLETS)}]\s*")
_LIGATURES = {"ﬁ": "fi", "ﬂ": "fl", "ﬀ": "ff", "ﬃ": "ffi", "ﬄ": "ffl"}
_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
# Lines that are just a page number, e.g. "2", "Page 2", "2 of 3"
_PAGE_NUM_RE = re.compile(r"^\s*(page\s*)?\d+(\s*(of|/)\s*\d+)?\s*$", re.I)


# --------------------------------------------------------------------------
# Cleaning
# --------------------------------------------------------------------------
def clean_text(raw: str) -> str:
    """Normalise raw extracted text without destroying line structure.

    Line breaks matter later (section headings and bullets are line-based),
    so we keep them and only fix obvious extraction artefacts.
    """
    if not raw:
        return ""

    # NFKC would turn ligatures into letters too, but we also map them explicitly.
    for lig, rep in _LIGATURES.items():
        raw = raw.replace(lig, rep)
    text = unicodedata.normalize("NFKC", raw)
    text = text.replace("\u00a0", " ").replace("\u200b", "")
    text = _CONTROL_RE.sub("", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # PDF fonts without a Unicode map come out as "(cid:127)". At the start of
    # a line this is almost always a bullet glyph; elsewhere it is just noise.
    text = re.sub(r"(?m)^\s*\(cid:\d+\)\s*", "• ", text)
    text = re.sub(r"\(cid:\d+\)", "", text)

    # Re-join words split by a hyphen at a line end: "develop-\nment" -> "development"
    text = re.sub(r"([a-z])-\n([a-z])", r"\1\2", text)

    lines: List[str] = []
    for line in text.split("\n"):
        line = re.sub(r"[ \t]+", " ", line).strip()
        if _PAGE_NUM_RE.match(line):
            continue
        # Normalise every bullet style to "- "
        if _BULLET_RE.match(line):
            line = "- " + _BULLET_RE.sub("", line, count=1)
        lines.append(line)

    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text)   # at most one blank line in a row
    return text.strip()


# --------------------------------------------------------------------------
# Format-specific readers: each returns (raw_text, n_pages, warnings)
# --------------------------------------------------------------------------
def _read_pdf(path: Path) -> Tuple[str, int, List[str]]:
    warnings: List[str] = []
    pages: List[str] = []
    n_pages = 0

    try:
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            n_pages = len(pdf.pages)
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
    except Exception as exc:  # corrupt file, encrypted, etc.
        warnings.append(f"pdfplumber failed ({exc}); trying pypdf")

    # Fallback: pypdf, if pdfplumber failed or returned nothing.
    if not "".join(pages).strip():
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            n_pages = len(reader.pages)
            pages = [(p.extract_text() or "") for p in reader.pages]
        except Exception as exc:
            warnings.append(f"pypdf failed ({exc})")

    text = "\n".join(pages)
    if not text.strip():
        warnings.append("No text layer found: this looks like a scanned/image PDF. "
                        "OCR would be needed (not implemented in the prototype).")
    return text, n_pages, warnings


def _read_docx(path: Path) -> Tuple[str, int, List[str]]:
    from docx import Document

    warnings: List[str] = []
    doc = Document(str(path))
    parts: List[str] = [p.text for p in doc.paragraphs]

    # Many resume templates put content inside tables.
    n_table_cells = 0
    for table in doc.tables:
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:      # merged cells repeat
                    continue
                seen.add(id(cell._tc))
                if cell.text.strip():
                    parts.append(cell.text)
                    n_table_cells += 1
    if n_table_cells:
        warnings.append(f"Read {n_table_cells} table cells; reading order may differ "
                        "from the visual layout.")
    return "\n".join(parts), 1, warnings


def _read_txt(path: Path) -> Tuple[str, int, List[str]]:
    for enc in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return path.read_text(encoding=enc), 1, []
        except UnicodeDecodeError:
            continue
    return "", 1, ["Could not decode text file."]


_READERS = {".pdf": _read_pdf, ".docx": _read_docx, ".txt": _read_txt}


# --------------------------------------------------------------------------
# Public API
# --------------------------------------------------------------------------
def extract_text(path: str | Path) -> ExtractionResult:
    """Extract and clean text from a PDF, DOCX or TXT resume."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file type '{ext}'. Supported: {sorted(SUPPORTED)}")

    raw, n_pages, warnings = _READERS[ext](path)
    text = clean_text(raw)

    if text and len(text) < 200:
        warnings.append("Very little text extracted (<200 chars); resume may be "
                        "image-based or mostly graphics.")

    return ExtractionResult(
        text=text,
        source_path=str(path),
        source_type=ext.lstrip("."),
        n_pages=n_pages,
        n_chars=len(text),
        warnings=warnings,
    )


def extract_text_from_bytes(data: bytes, filename: str) -> ExtractionResult:
    """Helper for the Streamlit uploader (which gives us bytes, not a path)."""
    import tempfile
    suffix = Path(filename).suffix.lower()
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(data)
        tmp_path = Path(tmp.name)
    try:
        result = extract_text(tmp_path)
        result.source_path = filename
        return result
    finally:
        tmp_path.unlink(missing_ok=True)


def _main(argv: List[str]) -> int:
    if len(argv) != 2:
        print("usage: python -m resume_analyzer.extraction <resume.pdf|docx|txt>")
        return 2
    res = extract_text(argv[1])
    print(f"[{res.source_type}] pages={res.n_pages} chars={res.n_chars}")
    for w in res.warnings:
        print(f"WARNING: {w}")
    print("-" * 60)
    print(res.text)
    return 0 if res.ok else 1


if __name__ == "__main__":
    sys.exit(_main(sys.argv))
